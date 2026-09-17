from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook


DOCUMENT_TYPES = {
    "packing_list": "Packing List",
    "hazardous_request": "Hazardous Cargo Request",
    "imo_declaration": "IMO Dangerous Goods Declaration",
    "multimodal_form": "Multimodal Dangerous Goods Form",
}

FIELD_DEFINITIONS = [
    # Shipment
    {"key": "product_name", "label": "Product name", "group": "Shipment", "required": True},
    {"key": "batch_number", "label": "Batch number", "group": "Shipment"},
    {"key": "invoice_reference", "label": "Invoice number and date", "group": "Shipment"},
    {"key": "buyer_order_reference", "label": "Buyer order number and date", "group": "Shipment"},
    {"key": "booking_number", "label": "Booking / transport document number", "group": "Shipment"},
    {"key": "shipper_reference", "label": "Shipper reference", "group": "Shipment"},
    {"key": "freight_forwarder_reference", "label": "Freight forwarder reference", "group": "Shipment"},
    {"key": "carrier", "label": "Carrier", "group": "Shipment"},
    {"key": "vessel_voyage", "label": "Vessel / flight and voyage", "group": "Shipment"},
    {"key": "port_loading", "label": "Port of loading", "group": "Shipment", "required": True},
    {"key": "port_discharge", "label": "Port of discharge", "group": "Shipment", "required": True},
    {"key": "final_destination", "label": "Final destination", "group": "Shipment", "required": True},
    {"key": "country_origin", "label": "Country of origin", "group": "Shipment"},
    {"key": "country_destination", "label": "Country of final destination", "group": "Shipment"},
    # Parties
    {"key": "shipper_name", "label": "Shipper / exporter", "group": "Parties", "required": True},
    {"key": "shipper_address", "label": "Shipper address", "group": "Parties", "multiline": True},
    {"key": "shipper_email", "label": "Shipper email", "group": "Parties"},
    {"key": "buyer_name", "label": "Buyer name", "group": "Parties"},
    {"key": "buyer_address", "label": "Buyer address", "group": "Parties", "multiline": True},
    {"key": "consignee_name", "label": "Consignee name", "group": "Parties", "required": True},
    {"key": "consignee_address", "label": "Consignee address", "group": "Parties", "multiline": True},
    {"key": "consignee_email", "label": "Consignee email", "group": "Parties"},
    # Quantity and packing
    {"key": "net_quantity", "label": "Net quantity", "group": "Quantity & packing", "required": True},
    {"key": "gross_mass", "label": "Gross mass", "group": "Quantity & packing", "required": True},
    {"key": "package_count", "label": "Number of packages", "group": "Quantity & packing", "required": True},
    {"key": "outer_packing", "label": "Outer packing", "group": "Quantity & packing"},
    {"key": "inner_packing", "label": "Inner packing", "group": "Quantity & packing"},
    {"key": "drum_weight", "label": "Drum weight calculation", "group": "Quantity & packing"},
    {"key": "pallet_weight", "label": "Pallet weight calculation", "group": "Quantity & packing"},
    {"key": "material_weight", "label": "Material weight", "group": "Quantity & packing"},
    {"key": "full_drums", "label": "Full drums", "group": "Quantity & packing"},
    {"key": "partial_drum", "label": "Partial drum", "group": "Quantity & packing"},
    {"key": "container_number", "label": "Container number", "group": "Quantity & packing"},
    {"key": "seal_number", "label": "Seal number", "group": "Quantity & packing"},
    {"key": "container_size", "label": "Container size and type", "group": "Quantity & packing"},
    {"key": "tare_mass", "label": "Container tare mass", "group": "Quantity & packing"},
    {"key": "cube", "label": "Cube (m³)", "group": "Quantity & packing"},
    # Dangerous goods
    {"key": "proper_shipping_name", "label": "Proper shipping name", "group": "Dangerous goods", "required": True, "multiline": True},
    {"key": "technical_name", "label": "Technical name", "group": "Dangerous goods", "required": True, "multiline": True},
    {"key": "un_number", "label": "UN number", "group": "Dangerous goods", "required": True},
    {"key": "imo_class", "label": "IMO class", "group": "Dangerous goods", "required": True},
    {"key": "packing_group", "label": "Packing group", "group": "Dangerous goods", "required": True},
    {"key": "cargo_status", "label": "Cargo state", "group": "Dangerous goods"},
    {"key": "limited_quantity", "label": "Limited / excepted quantity", "group": "Dangerous goods"},
    {"key": "flash_point", "label": "Flash point", "group": "Dangerous goods", "required": True},
    {"key": "subsidiary_risk", "label": "Subsidiary risk", "group": "Dangerous goods"},
    {"key": "ems_code", "label": "EMS code", "group": "Dangerous goods", "required": True},
    {"key": "marine_pollutant", "label": "Marine pollutant", "group": "Dangerous goods", "required": True},
    {"key": "un_packaging_code", "label": "UN packaging approval code", "group": "Dangerous goods", "required": True},
    {"key": "storage_temperature", "label": "Storage temperature", "group": "Dangerous goods"},
    {"key": "emergency_contact", "label": "Emergency contact person and number", "group": "Dangerous goods", "required": True},
    {"key": "un_test_marked", "label": "Outer packing UN test marked", "group": "Dangerous goods"},
    {"key": "hazard_chemicals", "label": "Chemicals giving rise to hazard", "group": "Dangerous goods"},
    {"key": "liable_flammable_gases", "label": "Liable to evolve flammable gases / vapours", "group": "Dangerous goods"},
    {"key": "discharge_permission", "label": "Discharge permission", "group": "Dangerous goods"},
    {"key": "imo_label", "label": "IMO label", "group": "Dangerous goods"},
    {"key": "mfag_number", "label": "MFAG number", "group": "Dangerous goods"},
    {"key": "reefer_temperature", "label": "Reefer temperature", "group": "Dangerous goods"},
    {"key": "humidity", "label": "Humidity", "group": "Dangerous goods"},
    {"key": "ventilation", "label": "Ventilation", "group": "Dangerous goods"},
    {"key": "boiling_point", "label": "Boiling point", "group": "Dangerous goods"},
    {"key": "notes", "label": "Additional notes", "group": "Dangerous goods", "multiline": True},
    # Sign-off
    {"key": "declarant_name", "label": "Declarant name / status", "group": "Sign-off"},
    {"key": "place_date", "label": "Place and date", "group": "Sign-off"},
]


def _text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"[ \t]+", " ", str(value).replace("\r", "")).strip()


def _after(value: Any, label: str) -> str:
    text = _text(value)
    match = re.search(re.escape(label) + r"\s*:?\s*(.*)$", text, flags=re.IGNORECASE | re.DOTALL)
    return _text(match.group(1)) if match else text


def _strip_prefix(value: Any, *prefixes: str) -> str:
    text = _text(value)
    for prefix in prefixes:
        text = re.sub(rf"^\s*{re.escape(prefix)}\s*:?\s*", "", text, flags=re.IGNORECASE)
    return _text(text)


def _openpyxl_fields(content: bytes) -> tuple[str, dict[str, str]]:
    workbook = load_workbook(io.BytesIO(content), data_only=False)
    sheet = workbook[workbook.sheetnames[0]]
    marker = " ".join(_text(sheet.cell(row, column).value) for row in range(1, min(sheet.max_row, 12) + 1) for column in range(1, min(sheet.max_column, 8) + 1)).upper()
    cell = lambda coordinate: _text(sheet[coordinate].value)
    if "PACKING LIST" in marker:
        return "packing_list", {
            "shipper_name": cell("B3"),
            "shipper_address": "\n".join(filter(None, [cell("B4"), cell("B5"), cell("B6")])),
            "invoice_reference": cell("F3"),
            "buyer_order_reference": cell("F5"),
            "buyer_name": cell("F10"),
            "buyer_address": "\n".join(filter(None, [cell("F11"), cell("F12")])),
            "consignee_name": cell("A12"),
            "consignee_address": "\n".join(filter(None, [cell("A13"), cell("A14")])),
            "country_origin": cell("F17"),
            "country_destination": cell("I17"),
            "port_loading": cell("C22"),
            "port_discharge": cell("A24"),
            "final_destination": cell("C24"),
            "net_quantity": _strip_prefix(cell("F23"), "Net weight"),
            "gross_mass": cell("J23"),
            "product_name": cell("E28"),
            "batch_number": _strip_prefix(cell("E29"), "BATCH NO"),
            "package_count": re.sub(r"\D+", "", cell("E30")) or cell("E30"),
            "outer_packing": cell("E30"),
            "drum_weight": _strip_prefix(cell("A28"), "Weight of Drums"),
            "pallet_weight": _strip_prefix(cell("A29"), "Weight of Pallet"),
            "material_weight": _strip_prefix(cell("A31"), "Weight of Material"),
            "full_drums": _strip_prefix(cell("E31"), "Full Drums"),
            "partial_drum": _strip_prefix(cell("E32"), "Partial Drum"),
        }
    if "IMO DANGEROUS GOODS DECLARATION" in marker:
        block = lambda row: _text(sheet[f"B{row}"].value)
        return "imo_declaration", {
            "shipper_name": "AROMAZEN PRIVATE LIMITED" if "AROMAZEN" in cell("A8").upper() else "",
            "shipper_address": _text(cell("A8")).replace("Shipper:", "").replace("AROMAZEN PRIVATE LIMITED", "").strip(),
            "shipper_email": cell("A9"),
            "consignee_name": _after(cell("A11").split("\n", 1)[0], "Consignee"),
            "consignee_address": cell("A11").split("\n", 1)[1] if "\n" in cell("A11") else "",
            "consignee_email": cell("A12"),
            "carrier": cell("E11"),
            "container_number": cell("E22"),
            "port_loading": cell("C22"),
            "port_discharge": cell("C23"),
            "proper_shipping_name": _strip_prefix(block(30), "Proper Shipping Name"),
            "technical_name": _strip_prefix(block(31), "Technical Name"),
            "cargo_status": _strip_prefix(block(32), "Properties"),
            "un_number": _strip_prefix(block(33), "UNNO", "UN"),
            "imo_class": _strip_prefix(block(34), "CLASS"),
            "ems_code": _strip_prefix(block(35), "EMS No"),
            "packing_group": _strip_prefix(block(36), "PKG Group"),
            "marine_pollutant": _strip_prefix(block(37), "Marine Pollutant"),
            "flash_point": _strip_prefix(block(38), "Flash point"),
            "net_quantity": _strip_prefix(block(39), "NT.WT"),
            "gross_mass": _strip_prefix(block(40), "GR.WT"),
            "package_count": _strip_prefix(block(42), "No of Packages"),
            "inner_packing": _strip_prefix(block(43), "INNER PACKING"),
            "outer_packing": _strip_prefix(block(44), "OUTER PKG"),
            "un_packaging_code": _strip_prefix(block(45), "UN PKG CODE"),
            "storage_temperature": _strip_prefix(block(46), "Storage temp"),
            "emergency_contact": _strip_prefix(block(47), "Emergency contact no / person"),
            "declarant_name": cell("E59"),
            "place_date": cell("E61"),
        }
    # Converted legacy Hazardous Cargo Request.
    if "HAZARDOUS CARGO REQUEST" in marker:
        return _hazardous_fields(lambda row, column=3: _text(sheet.cell(row, column).value))
    raise ValueError("The uploaded workbook is not one of the approved Merchandising document formats.")


def _hazardous_fields(value) -> tuple[str, dict[str, str]]:
    return "hazardous_request", {
        "product_name": value(2), "vessel_voyage": value(4), "container_size": value(5),
        "port_loading": value(6), "port_discharge": value(7), "final_destination": value(8),
        "booking_number": value(9), "outer_packing": value(10), "inner_packing": value(11),
        "un_test_marked": value(12), "un_number": _strip_prefix(value(13), "UN"),
        "imo_class": value(14), "proper_shipping_name": value(15), "technical_name": value(16),
        "cargo_status": value(17), "limited_quantity": value(18), "packing_group": value(19),
        "hazard_chemicals": value(21), "gross_mass": value(22), "net_quantity": value(23),
        "flash_point": value(24), "subsidiary_risk": value(25), "ems_code": value(26),
        "marine_pollutant": value(27), "liable_flammable_gases": value(29),
        "discharge_permission": value(30), "shipper_name": value(31), "consignee_name": value(32),
        "emergency_contact": value(33), "un_packaging_code": value(34),
    }


def _xls_fields(content: bytes) -> tuple[str, dict[str, str]]:
    try:
        import xlrd  # type: ignore
    except ImportError as exc:  # pragma: no cover - dependency is installed in production
        raise ValueError("Legacy XLS reading is unavailable on this server. Save the file as XLSX and try again.") from exc
    book = xlrd.open_workbook(file_contents=content)
    sheet = book.sheet_by_index(0)
    marker = " ".join(_text(sheet.cell_value(row, column)) for row in range(min(sheet.nrows, 12)) for column in range(min(sheet.ncols, 7))).upper()
    if "HAZARDOUS CARGO REQUEST" not in marker:
        raise ValueError("The uploaded XLS file is not the approved Hazardous Cargo Request format.")
    return _hazardous_fields(lambda row, column=3: _text(sheet.cell_value(row - 1, column - 1)) if row <= sheet.nrows and column <= sheet.ncols else "")


def _docx_fields(content: bytes) -> tuple[str, dict[str, str]]:
    document = Document(io.BytesIO(content))
    if not document.tables:
        raise ValueError("The uploaded Word file does not contain the approved Multimodal form table.")
    table = document.tables[0]
    value = lambda row, column: _text(table.rows[row].cells[column].text)
    goods = value(10, 0)
    def capture(label: str, next_label: str | None = None) -> str:
        end = rf"(?={re.escape(next_label)})" if next_label else "$"
        match = re.search(rf"{re.escape(label)}\s*:?\s*(.*?){end}", goods, flags=re.IGNORECASE | re.DOTALL)
        return _text(match.group(1)) if match else ""
    return "multimodal_form", {
        "shipper_name": "AROMAZEN PRIVATE LIMITED" if "AROMAZEN" in value(0, 0).upper() else "",
        "booking_number": _after(value(0, 2), "2 Transport document number"),
        "shipper_reference": _after(value(1, 3), "4 Shipper's reference"),
        "freight_forwarder_reference": _after(value(2, 3), "5 Freight Forwarder's reference"),
        "consignee_name": "Creative Lights Vietnam HD co., Ltd" if "CREATIVE LIGHTS" in value(3, 0).upper() else "",
        "carrier": _after(value(3, 2), "Carrier"),
        "vessel_voyage": _after(value(7, 0), "10 Vessel and Voyage number"),
        "port_loading": _after(value(7, 1), "11 Port of loading"),
        "port_discharge": _after(value(8, 0), "12 Port of discharge"),
        "final_destination": _after(value(8, 1), "13 Destination"),
        "un_number": _strip_prefix(capture("UN NO", "Proper Shipping Name"), "UN"),
        "proper_shipping_name": capture("Proper Shipping Name", "Technical Name /chemical name"),
        "technical_name": capture("Technical Name /chemical name", "Class"),
        "imo_class": capture("Class", "Sub Risk"),
        "subsidiary_risk": capture("Sub Risk (if any)", "Packing Group"),
        "packing_group": capture("Packing Group", "Marine Pollutant"),
        "marine_pollutant": capture("Marine Pollutant", "Packaging type"),
        "outer_packing": capture("Packaging type (Outer) with quantity", "Inner packing details"),
        "inner_packing": capture("Inner packing details", "Flashpoint"),
        "flash_point": capture("Flashpoint", "EMS Code"),
        "ems_code": capture("EMS Code", "IMO Label"),
        "imo_label": capture("IMO Label", "MFAG Number"),
        "mfag_number": capture("MFAG Number", "Reefer Temp"),
        "emergency_contact": capture("Emergency Contact person name/number", "Limited Quantity"),
        "limited_quantity": capture("Limited Quantity", "Poisonous inhalation hazard"),
        "gross_mass": value(10, 3), "net_quantity": value(10, 4), "cube": value(10, 5),
        "un_packaging_code": _after(value(12, 0), "Other Details UN PACKAGING CODE -"),
        "container_number": _after(value(13, 0), "15 Container identification no."),
        "seal_number": _after(value(13, 1), "16 Seal number(s)"),
        "container_size": _after(value(13, 2), "17 Container size & type"),
        "tare_mass": _after(value(13, 3), "18 Tare mass (kg)"),
        "declarant_name": _after(value(16, 0), "Name/status of declarant"),
        "place_date": _after(value(17, 0), "Place and date"),
    }


def extract_fields(content: bytes, filename: str) -> tuple[str, dict[str, str]]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        document_type, values = _openpyxl_fields(content)
    elif suffix == ".xls":
        document_type, values = _xls_fields(content)
    elif suffix == ".docx":
        document_type, values = _docx_fields(content)
    else:
        raise ValueError("Choose an approved XLS, XLSX, XLSM, or DOCX Merchandising document.")
    return document_type, {key: _text(value) for key, value in values.items() if _text(value)}


def missing_required(fields: dict[str, str]) -> list[str]:
    return [item["label"] for item in FIELD_DEFINITIONS if item.get("required") and not _text(fields.get(item["key"]))]


def _set(sheet, coordinate: str, value: Any) -> None:
    sheet[coordinate] = _text(value)


def _generate_packing_list(template: Path, output: Path, fields: dict[str, str]) -> None:
    workbook = load_workbook(template)
    sheet = workbook[workbook.sheetnames[0]]
    address = [line.strip() for line in fields.get("shipper_address", "").splitlines() if line.strip()]
    consignee = [line.strip() for line in fields.get("consignee_address", "").splitlines() if line.strip()]
    buyer = [line.strip() for line in fields.get("buyer_address", "").splitlines() if line.strip()]
    values = {
        "B3": fields.get("shipper_name"), "B4": address[0] if address else "", "B5": address[1] if len(address) > 1 else "", "B6": address[2] if len(address) > 2 else "",
        "F3": fields.get("invoice_reference"), "F5": fields.get("buyer_order_reference"),
        "F10": fields.get("buyer_name"), "F11": buyer[0] if buyer else "", "F12": "\n".join(buyer[1:]),
        "A12": fields.get("consignee_name"), "A13": consignee[0] if consignee else "", "A14": "\n".join(consignee[1:]),
        "F17": fields.get("country_origin"), "I17": fields.get("country_destination"),
        "A22": fields.get("vessel_voyage"), "C22": fields.get("port_loading"), "A24": fields.get("port_discharge"), "C24": fields.get("final_destination"),
        "F23": f"Net weight: {fields.get('net_quantity', '')}".strip(), "J23": fields.get("gross_mass"), "I27": fields.get("net_quantity"),
        "A28": f"Weight of Drums : {fields.get('drum_weight', '')}".strip(), "A29": f"Weight of Pallet : {fields.get('pallet_weight', '')}".strip(),
        "A31": f"Weight of Material : {fields.get('material_weight') or fields.get('net_quantity', '')}".strip(), "A32": f"Total: {fields.get('gross_mass', '')}".strip(),
        "A33": f"Total Drums : {fields.get('package_count', '')}".strip(), "E28": fields.get("product_name"),
        "E29": f"BATCH NO: {fields.get('batch_number', '')}".strip(), "E30": fields.get("outer_packing") or f"DRUMS: {fields.get('package_count', '')}",
        "E31": f"Full Drums = {fields.get('full_drums', '')}".strip(), "E32": f"Partial Drum = {fields.get('partial_drum', '')}".strip(),
    }
    for coordinate, value in values.items():
        _set(sheet, coordinate, value)
    sheet.print_area = "A1:K50"
    workbook.save(output)


def _generate_hazardous_request(template: Path, output: Path, fields: dict[str, str]) -> None:
    workbook = load_workbook(template)
    sheet = workbook[workbook.sheetnames[0]]
    row_values = {
        2: fields.get("product_name"), 4: fields.get("vessel_voyage"), 5: fields.get("container_size"), 6: fields.get("port_loading"),
        7: fields.get("port_discharge"), 8: fields.get("final_destination"), 9: fields.get("booking_number"), 10: fields.get("outer_packing"),
        11: fields.get("inner_packing"), 12: fields.get("un_test_marked"), 13: fields.get("un_number"), 14: fields.get("imo_class"),
        15: fields.get("proper_shipping_name"), 16: fields.get("technical_name"), 17: fields.get("cargo_status"), 18: fields.get("limited_quantity"),
        19: fields.get("packing_group"), 21: fields.get("hazard_chemicals"), 22: fields.get("gross_mass"), 23: fields.get("net_quantity"),
        24: fields.get("flash_point"), 25: fields.get("subsidiary_risk"), 26: fields.get("ems_code"), 27: fields.get("marine_pollutant"),
        29: fields.get("liable_flammable_gases"), 30: fields.get("discharge_permission"), 31: fields.get("shipper_name"), 32: fields.get("consignee_name"),
        33: fields.get("emergency_contact"), 34: fields.get("un_packaging_code"),
    }
    for row, value in row_values.items():
        sheet.cell(row, 3).value = _text(value)
    sheet.print_area = "A1:G40"
    workbook.save(output)


def _generate_imo_declaration(template: Path, output: Path, fields: dict[str, str]) -> None:
    workbook = load_workbook(template)
    sheet = workbook[workbook.sheetnames[0]]
    values = {
        "A8": f"Shipper: {fields.get('shipper_name', '')}\n{fields.get('shipper_address', '')}", "A9": fields.get("shipper_email"),
        "A11": f"Consignee : {fields.get('consignee_name', '')}\n{fields.get('consignee_address', '')}", "A12": fields.get("consignee_email"),
        "E8": fields.get("booking_number"), "E11": fields.get("carrier"), "E15": fields.get("shipper_name"),
        "E17": fields.get("place_date"), "E21": fields.get("container_number"), "C22": fields.get("port_loading"), "C23": fields.get("port_discharge"),
        "B30": f"Proper Shipping Name: {fields.get('proper_shipping_name', '')}", "B31": f"Technical Name: {fields.get('technical_name', '')}",
        "B32": f"Properties : {fields.get('cargo_status', '')}", "B33": f"UNNO: {fields.get('un_number', '')}", "B34": f"CLASS: {fields.get('imo_class', '')}",
        "B35": f"EMS No: {fields.get('ems_code', '')}", "B36": f"PKG Group: {fields.get('packing_group', '')}",
        "B37": f"Marine Pollutant: {fields.get('marine_pollutant', '')}", "B38": f"Flash point: {fields.get('flash_point', '')}",
        "B39": f"NT.WT: {fields.get('net_quantity', '')}", "B40": f"GR.WT: {fields.get('gross_mass', '')}",
        "B41": f"Cargo status-Solid/Liquid/Gas/Paste : {fields.get('cargo_status', '')}", "B42": f"No of Packages: {fields.get('package_count', '')}",
        "B43": f"INNER PACKING : {fields.get('inner_packing', '')}", "B44": f"OUTER PKG : {fields.get('outer_packing', '')}",
        "B45": f"UN PKG CODE : {fields.get('un_packaging_code', '')}", "B46": f"Storage temp : {fields.get('storage_temperature', '')}",
        "B47": f"Emergency contact no / person : {fields.get('emergency_contact', '')}", "E59": fields.get("shipper_name"), "E61": fields.get("place_date"),
    }
    for coordinate, value in values.items():
        _set(sheet, coordinate, value)
    # The supplied workbook contains drawing anchors outside its visible form.
    # An explicit print area prevents those objects from shrinking the form to
    # a thumbnail when Excel exports or prints the generated workbook.
    sheet.print_area = "A1:H64"
    workbook.save(output)


def _set_cell(table, row: int, column: int, value: str) -> None:
    table.rows[row].cells[column].text = value


def _generate_multimodal(template: Path, output: Path, fields: dict[str, str]) -> None:
    document = Document(template)
    table = document.tables[0]
    _set_cell(table, 0, 0, f"1 Shipper/Consignor/Sender (full style address is mandatory)\n\n{fields.get('shipper_name', '')}\n{fields.get('shipper_address', '')}")
    _set_cell(table, 0, 2, f"2 Transport document number\n{fields.get('booking_number', '')}")
    _set_cell(table, 1, 3, f"4 Shipper's reference\n{fields.get('shipper_reference', '')}")
    _set_cell(table, 2, 3, f"5 Freight Forwarder's reference\n{fields.get('freight_forwarder_reference', '')}")
    _set_cell(table, 3, 0, f"6 Consignee\n{fields.get('consignee_name', '')}\n{fields.get('consignee_address', '')}")
    _set_cell(table, 3, 2, f"Carrier:\n{fields.get('carrier', '')}")
    _set_cell(table, 7, 0, f"10 Vessel and Voyage number\n{fields.get('vessel_voyage', '')}")
    _set_cell(table, 7, 1, f"11 Port of loading\n{fields.get('port_loading', '')}")
    _set_cell(table, 8, 0, f"12 Port of discharge\n{fields.get('port_discharge', '')}")
    _set_cell(table, 8, 1, f"13 Destination\n{fields.get('final_destination', '')}")
    goods = (
        f"UN NO: {fields.get('un_number', '')}\n"
        f"Proper Shipping Name: {fields.get('proper_shipping_name', '')}\n"
        f"Technical Name /chemical name: {fields.get('technical_name', '')}\n"
        f"Class: {fields.get('imo_class', '')}\nSub Risk (if any): {fields.get('subsidiary_risk', '')}\n"
        f"Packing Group: {fields.get('packing_group', '')}\nMarine Pollutant: {fields.get('marine_pollutant', '')}\n"
        f"Packaging type (Outer) with quantity: {fields.get('outer_packing', '')}\nInner packing details: {fields.get('inner_packing', '')}\n"
        f"Flashpoint: {fields.get('flash_point', '')}\nEMS Code: {fields.get('ems_code', '')}\n"
        f"IMO Label: {fields.get('imo_label', '')}\nMFAG Number: {fields.get('mfag_number', '')}\n"
        f"Reefer Temp: {fields.get('reefer_temperature', '')}    Humidity: {fields.get('humidity', '')}    Ventilation: {fields.get('ventilation', '')}\n"
        f"Boiling Point: {fields.get('boiling_point', '')}\nEmergency Contact person name/number: {fields.get('emergency_contact', '')}\n"
        f"Limited Quantity: {fields.get('limited_quantity', '')}\nNotes: {fields.get('notes', '')}"
    )
    _set_cell(table, 10, 0, goods)
    _set_cell(table, 10, 3, fields.get("gross_mass", "")); _set_cell(table, 10, 4, fields.get("net_quantity", "")); _set_cell(table, 10, 5, fields.get("cube", ""))
    _set_cell(table, 12, 0, f"Other Details      UN PACKAGING CODE - {fields.get('un_packaging_code', '')}")
    _set_cell(table, 13, 0, f"15 Container identification no.\n{fields.get('container_number', '')}")
    _set_cell(table, 13, 1, f"16 Seal number(s)\n{fields.get('seal_number', '')}")
    _set_cell(table, 13, 2, f"17 Container size & type\n{fields.get('container_size', '')}")
    _set_cell(table, 13, 3, f"18 Tare mass (kg)\n{fields.get('tare_mass', '')}")
    _set_cell(table, 13, 4, f"19 Total gross (incl. tare) (kg)\n{fields.get('gross_mass', '')}")
    _set_cell(table, 15, 0, f"20 Name of company (stamp & signature are mandatory)\n{fields.get('shipper_name', '')}")
    _set_cell(table, 16, 0, f"Name/status of declarant\n{fields.get('declarant_name', '')}")
    _set_cell(table, 17, 0, f"Place and date\n{fields.get('place_date', '')}")
    document.save(output)


def generate_document(template: Path, output: Path, document_type: str, fields: dict[str, str]) -> None:
    if document_type == "packing_list":
        _generate_packing_list(template, output, fields)
    elif document_type == "hazardous_request":
        _generate_hazardous_request(template, output, fields)
    elif document_type == "imo_declaration":
        _generate_imo_declaration(template, output, fields)
    elif document_type == "multimodal_form":
        _generate_multimodal(template, output, fields)
    else:
        raise ValueError("Unsupported Merchandising document type.")
