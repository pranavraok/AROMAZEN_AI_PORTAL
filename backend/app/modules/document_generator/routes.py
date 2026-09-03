import io
import json
import re
import time
import uuid
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import Workbook
from pydantic import BaseModel, Field
from sqlalchemy import or_, select
from sqlalchemy.ext.asyncio import AsyncSession
import httpx
from docx import Document as WordDocument

from app.core.config import get_settings
from app.db.session import get_db_session
from app.modules.ai.providers import AIProviderRouter, ProviderError, estimate_cost
from app.modules.document_generator.engine import coa_parameter_rows, field_schema, generate_docx, normalise, read_excel
from app.modules.identity.authorization import require_department, require_permissions
from app.modules.identity.models import AIUsageEvent, AuditEvent, Department, DocumentGeneration, KnowledgeCollection, KnowledgeDocument, User, collection_departments
from app.modules.settings.service import provider_runtime_settings
from app.modules.identity.service import role_keys_for_user
from app.modules.knowledge.storage import organized_storage_name
from app.modules.knowledge.department_uploads import DepartmentUpload, replace_department_uploads
from app.modules.hr_letters.routes import _convert_docx_to_pdf

router = APIRouter(dependencies=[Depends(require_department("r-d"))])


class DraftNotesRequest(BaseModel):
    template_document_id: str
    notes: str = Field(min_length=1, max_length=20000)
    current_fields: dict[str, str] = Field(default_factory=dict)
    current_rows: list[dict[str, str]] = Field(default_factory=list, max_length=200)
    field_labels: dict[str, str] = Field(default_factory=dict)


def _parse_ai_json(text: str) -> dict:
    start, end = text.find("{"), text.rfind("}")
    if start < 0 or end <= start:
        raise ValueError("No JSON object was returned.")
    value = json.loads(text[start:end + 1])
    if not isinstance(value, dict):
        raise ValueError("The AI response was not an object.")
    return value


def _quick_field_updates(notes: str, document_type: str, allowed_fields: set[str]) -> dict[str, str]:
    """Instantly map common spoken labels before spending time or credits on an AI call."""
    coa_aliases = {
        "name of the product": "product_name", "name of product": "product_name", "product name": "product_name", "name": "product_name",
        "product code": "product_code", "code": "product_code", "batch number": "batch_number", "batch no": "batch_number", "batch": "batch_number",
        "customer name": "customer_name", "customer": "customer_name", "date of manufacturing": "manufacturing_date", "manufacturing date": "manufacturing_date", "manufacturing": "manufacturing_date", "mfg date": "manufacturing_date",
        "expiry date": "expiry_date", "expiration date": "expiry_date", "expiry": "expiry_date", "expiration": "expiry_date", "quantity": "quantity", "storage condition": "storage_condition", "storage": "storage_condition",
        "tested by": "tested_by", "checked by": "checked_by", "date": "date",
    }
    sds_aliases = {
        "product identifier": "product_identifier", "product name": "product_identifier", "name": "product_identifier",
        "other identifiers": "other_identifiers", "recommended use": "recommended_use", "product use": "recommended_use", "use": "recommended_use",
        "supplier name": "supplier_name", "company name": "supplier_name", "supplier address": "supplier_address", "company address": "supplier_address",
        "supplier contact": "supplier_contact", "contact": "supplier_contact", "supplier phone": "supplier_phone", "company phone": "supplier_phone",
        "emergency phone": "emergency_phone", "hazard classification": "classification", "classification": "classification", "signal word": "signal_word",
        "hazard statements": "hazard_statements", "precautionary statements": "precautionary_statements", "other hazards": "other_hazards",
        "appearance": "appearance", "colour": "colour", "color": "colour", "odour": "odour", "odor": "odour", "ph": "ph",
        "flash point": "flash_point", "solubility": "solubility", "vapour pressure": "vapour_pressure", "vapor pressure": "vapour_pressure",
        "relative density": "relative_density", "revision date": "revision_date",
    }
    aliases = coa_aliases if document_type == "coa" else sds_aliases
    # COA test labels are boundaries too. This prevents a value such as Quantity
    # from swallowing all test results that follow it in one spoken sentence.
    boundaries = ({"appearance", "odour", "odor", "specific gravity", "flash point", "fire point", "refractive index"} if document_type == "coa" else set())
    usable = {alias: key for alias, key in aliases.items() if key in allowed_fields}
    all_aliases = set(usable).union(boundaries)
    alias_pattern = "|".join(re.escape(alias) for alias in sorted(all_aliases, key=len, reverse=True))
    if not alias_pattern:
        return {}
    matches = list(re.finditer(rf"(?i)\b({alias_pattern})\b", notes))
    updates: dict[str, str] = {}
    for index, match in enumerate(matches):
        key = usable.get(match.group(1).lower())
        if not key:
            continue
        end = matches[index + 1].start() if index + 1 < len(matches) else len(notes)
        value = notes[match.end():end]
        value = re.sub(r"^\s*(?:is|equals|equal to|=|:|as)\s*", "", value, flags=re.IGNORECASE).strip(" \t\r\n,;.")
        correction = re.split(r"(?i)\b(?:sorry(?:\s*,)?\s*)?(?:no(?:\s*,)?\s*)?(?:please\s+)?(?:correct\s+(?:that|it)\s+to|change\s+(?:that|it)\s+to|make\s+that|i\s+mean)\b|\bsorry\b", value)
        if len(correction) > 1 and correction[-1].strip():
            value = correction[-1].strip(" \t\r\n,;.")
        if value:
            updates[key] = _normalize_field_value(key, value)
    return updates


def _professional_date(value: str) -> str:
    """Standardize spoken Indian/English dates as '28 July 2026'."""
    cleaned = value.strip(" \t\r\n,;.")
    months = {name.lower(): index for index, name in enumerate(("January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"), 1)}
    month_pattern = "|".join(months)
    match = re.fullmatch(rf"(?i)(\d{{1,2}})(?:st|nd|rd|th)?(?:\s+of)?\s+({month_pattern})(?:\s*,?\s*(\d{{4}}))?", cleaned)
    if match:
        day, month_name, year = int(match.group(1)), match.group(2).lower(), int(match.group(3) or datetime.now().year)
        try:
            parsed = datetime(year, months[month_name], day)
            return f"{parsed.day} {parsed.strftime('%B')} {parsed.year}"
        except (ValueError, OSError):
            return f"{day} {datetime(year, months[month_name], 1).strftime('%B')} {year}"
    match = re.fullmatch(rf"(?i)({month_pattern})\s+(\d{{1,2}})(?:st|nd|rd|th)?(?:\s*,?\s*(\d{{4}}))?", cleaned)
    if match:
        month_name, day, year = match.group(1).lower(), int(match.group(2)), int(match.group(3) or datetime.now().year)
        try:
            parsed = datetime(year, months[month_name], day)
            return f"{parsed.day} {parsed.strftime('%B')} {parsed.year}"
        except (ValueError, OSError):
            return f"{day} {datetime(year, months[month_name], 1).strftime('%B')} {year}"
    match = re.fullmatch(r"(\d{1,2})[./-](\d{1,2})[./-](\d{4})", cleaned)
    if match:
        try:
            parsed = datetime(int(match.group(3)), int(match.group(2)), int(match.group(1)))
            return f"{parsed.day} {parsed.strftime('%B')} {parsed.year}"
        except (ValueError, OSError):
            return cleaned
    # Speech recognition commonly compacts Indian dates: "14 7" -> "147"
    # and "13 7 2028" -> "137 2028".
    match = re.fullmatch(r"(\d{2})(\d)(?:\s+(\d{4}))?", cleaned)
    if match:
        day, month, year = int(match.group(1)), int(match.group(2)), int(match.group(3) or datetime.now().year)
        try:
            parsed = datetime(year, month, day)
            return f"{parsed.day} {parsed.strftime('%B')} {parsed.year}"
        except (ValueError, OSError):
            pass
    match = re.fullmatch(r"(\d{1,2})\s+(\d{1,2})(?:\s+(\d{4}))?", cleaned)
    if match:
        day, month, year = int(match.group(1)), int(match.group(2)), int(match.group(3) or datetime.now().year)
        try:
            parsed = datetime(year, month, day)
            return f"{parsed.day} {parsed.strftime('%B')} {parsed.year}"
        except (ValueError, OSError):
            pass
    return cleaned


def _normalize_field_value(key: str, value: str) -> str:
    value = re.sub(r"\s+", " ", value).strip(" \t\r\n,;.")
    value = re.sub(r"(?i)\s+and$", "", value).strip()
    if key in {"date", "manufacturing_date", "expiry_date", "revision_date"}:
        return _professional_date(value)
    if key == "product_code":
        compact = re.sub(r"[^A-Za-z0-9]", "", value).upper()
        match = re.fullmatch(r"([A-Z]+)(\d+)", compact)
        return f"{match.group(1)} {match.group(2)}" if match else value.upper()
    if key == "batch_number":
        return re.sub(r"\s+", "", value).upper()
    if key == "quantity":
        match = re.fullmatch(r"(?i)(\d+(?:\.\d+)?)\s*(kg|kilograms?|g|grams?|mg|milligrams?|l|ml|millilitres?|milliliters?|litres?|liters?)", value)
        if match:
            unit = {
                "KILOGRAM": "KG", "KILOGRAMS": "KG", "GRAM": "G", "GRAMS": "G",
                "MILLIGRAM": "MG", "MILLIGRAMS": "MG", "LITRE": "L", "LITRES": "L",
                "LITER": "L", "LITERS": "L", "MILLILITRE": "ML", "MILLILITRES": "ML",
                "MILLILITER": "ML", "MILLILITERS": "ML",
            }.get(match.group(2).upper(), match.group(2).upper())
            return f"{match.group(1)} {unit}"
    if key in {"product_name", "customer_name", "tested_by", "checked_by", "supplier_name"}:
        return value.title()
    return value[:4000]


def _normalize_coa_value(parameter: str, column: str, value: str) -> str:
    """Apply professional COA notation and repair common speech artifacts."""
    text = re.sub(r"\s+", " ", str(value)).strip(" \t\r\n,;.")
    if not text:
        return ""
    # Spoken ranges must become a visible range, never concatenated decimals.
    text = re.sub(r"(?i)(\d(?:\.\d+)?)\s+(?:to|through)\s+(\d(?:\.\d+)?)", r"\1-\2", text)
    spoken_to_range = re.fullmatch(r"(\d+\.\d{2})2(\d+\.\d{2})", text)
    collapsed_range = re.fullmatch(r"(\d+\.\d{2})(\d+\.\d{2})", text)
    if spoken_to_range:
        text = f"{spoken_to_range.group(1)}-{spoken_to_range.group(2)}"
    elif collapsed_range:
        text = f"{collapsed_range.group(1)}-{collapsed_range.group(2)}"
    text = re.sub(r"(?i)\s*degrees?\s*(?:celsius|centigrade|c)\b", "°C", text)
    if column == "result" and normalise(text) in {"pass", "passes", "passed"}:
        return "Passes"
    if column == "specification" and normalise(text) == "for record":
        return "For record"
    # "Pale" is a frequent COA transcription error (fail/tale) for Appearance.
    if normalise(parameter) == "appearance" and column == "specification":
        text = re.sub(r"(?i)^(?:fail|tale)\s*,?\s*", "Pale ", text)
    return text[:2000]


DOCUMENT_DEPARTMENT_SLUGS = {"r-d", "qa", "quality-assurance", "qa-qc", "qa-and-qc", "quality-assurance-quality-control"}
QA_DEPARTMENT_SLUGS = {"qa", "quality-assurance", "qa-qc", "qa-and-qc", "quality-assurance-quality-control"}
QA_COA_MASTER_SOURCE = "qa-coa-master"
QA_COA_CANVA_URL = "https://www.canva.com/d/22DzkdhTpOfj6CV"


async def _require_document_department(session: AsyncSession, user: User) -> str:
    roles = await role_keys_for_user(session, user.id)
    if roles.intersection({"owner", "super_admin"}):
        return "r-d"
    department = await session.get(Department, user.department_id) if user.department_id else None
    if not department or department.slug not in DOCUMENT_DEPARTMENT_SLUGS:
        raise HTTPException(status_code=403, detail="SDS and COA creation is limited to the R&D and Quality Assurance departments.")
    return department.slug


async def _require_qa_department(session: AsyncSession, user: User) -> None:
    roles = await role_keys_for_user(session, user.id)
    if roles.intersection({"owner", "super_admin"}):
        return
    department = await session.get(Department, user.department_id) if user.department_id else None
    if not department or department.slug not in QA_DEPARTMENT_SLUGS:
        raise HTTPException(status_code=403, detail="This master template is restricted to the Quality Assurance department.")


def _type_for(name: str, source_key: str | None = None) -> str:
    if source_key and source_key.startswith("document-generator-template:"):
        template_type = source_key.split(":", 2)[1]
        if template_type in {"coa", "sds"}:
            return template_type
    return "sds" if "sds" in name.lower() else "coa"


async def _template(session: AsyncSession, user: User, template_id: str) -> tuple[KnowledgeDocument, KnowledgeCollection]:
    document = await session.get(KnowledgeDocument, template_id)
    collection = await session.get(KnowledgeCollection, document.collection_id) if document else None
    is_another_users_personal_template = bool(
        document
        and document.source_key
        and document.source_key.startswith("document-generator-template:")
        and document.source_key.count(":") >= 2
        and document.uploaded_by_user_id != user.id
    )
    belongs_to_document_department = bool(await session.scalar(select(collection_departments.c.collection_id).join(
        Department, Department.id == collection_departments.c.department_id
    ).where(
        collection_departments.c.collection_id == collection.id,
        Department.slug.in_(tuple(DOCUMENT_DEPARTMENT_SLUGS)),
    ))) if collection else False
    if (
        not document
        or not collection
        or document.organization_id != user.organization_id
        or document.status != "ready"
        or document.document_category != "document_template"
        or not belongs_to_document_department
        or is_another_users_personal_template
        or Path(document.original_filename).suffix.lower() != ".docx"
    ):
        raise HTTPException(status_code=404, detail="Word template not found.")
    return document, collection


@router.get("/templates")
async def list_templates(user: User = Depends(require_permissions("ai.workspace.use", "knowledge.read")), session: AsyncSession = Depends(get_db_session)) -> list[dict]:
    await _require_document_department(session, user)
    query = select(KnowledgeDocument, KnowledgeCollection).join(KnowledgeCollection, KnowledgeCollection.id == KnowledgeDocument.collection_id).join(
        collection_departments, collection_departments.c.collection_id == KnowledgeCollection.id
    ).join(Department, Department.id == collection_departments.c.department_id).where(
        KnowledgeDocument.organization_id == user.organization_id,
        KnowledgeDocument.status == "ready",
        KnowledgeDocument.document_category == "document_template",
        KnowledgeDocument.original_filename.ilike("%.docx"),
        KnowledgeCollection.status == "active",
        Department.slug.in_(tuple(DOCUMENT_DEPARTMENT_SLUGS)),
        or_(
            KnowledgeDocument.source_key.is_(None),
            KnowledgeDocument.source_key.not_like("document-generator-template:%:%"),
            KnowledgeDocument.uploaded_by_user_id == user.id,
        ),
    ).order_by(KnowledgeDocument.created_at.desc())
    result = []
    for document, collection in (await session.execute(query)).all():
        result.append({
            "id": str(document.id), "name": document.original_filename,
            "collection_name": collection.name, "document_type": _type_for(document.original_filename, document.source_key),
            "version": document.version, "source_key": document.source_key,
            "external_edit_url": document.external_edit_url,
        })
    return result


@router.post("/templates")
async def upload_template(
    document_type: str = Form(...),
    template_file: UploadFile = File(...),
    user: User = Depends(require_permissions("ai.workspace.use", "knowledge.read")),
    session: AsyncSession = Depends(get_db_session),
) -> dict:
    department_slug = await _require_document_department(session, user)
    if document_type not in {"coa", "sds"}:
        raise HTTPException(status_code=422, detail="Choose either an SDS or COA template type.")
    original_name = Path(template_file.filename or "").name
    if Path(original_name).suffix.lower() != ".docx":
        raise HTTPException(status_code=422, detail="Templates must be uploaded as DOCX Word files.")
    settings = get_settings()
    content = await template_file.read(settings.max_upload_size_mb * 1024 * 1024 + 1)
    if not content or len(content) > settings.max_upload_size_mb * 1024 * 1024:
        raise HTTPException(status_code=413, detail=f"The template is empty or exceeds the {settings.max_upload_size_mb} MB limit.")
    try:
        WordDocument(io.BytesIO(content))
    except Exception as exc:
        raise HTTPException(status_code=422, detail="The uploaded file is not a valid DOCX Word document.") from exc
    type_label = document_type.upper()
    stored_name = original_name if document_type in original_name.lower() else f"{Path(original_name).stem}-{type_label}.docx"
    documents = await replace_department_uploads(session, user, department_slug, [DepartmentUpload(
        f"document-generator-template:{document_type}:{user.id}",
        content,
        stored_name,
        template_file.content_type or "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        document_category="document_template",
    )])
    document = documents[0]
    collection = await session.get(KnowledgeCollection, document.collection_id)
    return {
        "id": str(document.id),
        "name": document.original_filename,
        "collection_name": collection.name if collection else "R&D",
        "document_type": document_type,
        "version": document.version,
        "source_key": document.source_key,
        "external_edit_url": document.external_edit_url,
    }


@router.post("/templates/coa-master")
async def replace_qa_coa_master(
    template_file: UploadFile = File(...),
    user: User = Depends(require_permissions("ai.workspace.use", "knowledge.read", "knowledge.write")),
    session: AsyncSession = Depends(get_db_session),
) -> dict:
    await _require_qa_department(session, user)
    original_name = Path(template_file.filename or "").name
    if Path(original_name).suffix.lower() != ".docx":
        raise HTTPException(status_code=422, detail="The COA master must be a DOCX Word file.")
    settings = get_settings()
    content = await template_file.read(settings.max_upload_size_mb * 1024 * 1024 + 1)
    if not content or len(content) > settings.max_upload_size_mb * 1024 * 1024:
        raise HTTPException(status_code=413, detail=f"The template is empty or exceeds the {settings.max_upload_size_mb} MB limit.")
    try:
        WordDocument(io.BytesIO(content))
    except Exception as exc:
        raise HTTPException(status_code=422, detail="The uploaded file is not a valid DOCX Word document.") from exc
    document = (await replace_department_uploads(session, user, "quality-assurance", [DepartmentUpload(
        QA_COA_MASTER_SOURCE,
        content,
        "AROMAZEN COA Master.docx",
        template_file.content_type or "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        document_category="document_template",
    )]))[0]
    document.external_edit_url = QA_COA_CANVA_URL
    await session.commit()
    collection = await session.get(KnowledgeCollection, document.collection_id)
    return {
        "id": str(document.id), "name": document.original_filename,
        "collection_name": collection.name if collection else "Quality Assurance",
        "document_type": "coa", "version": document.version,
        "source_key": document.source_key, "external_edit_url": document.external_edit_url,
    }


@router.get("/templates/{template_id}/content")
async def template_content(template_id: str, user: User = Depends(require_permissions("ai.workspace.use", "knowledge.read")), session: AsyncSession = Depends(get_db_session)) -> FileResponse:
    await _require_document_department(session, user)
    document, _ = await _template(session, user, template_id)
    path = Path(get_settings().upload_storage_path) / document.stored_filename
    if not path.is_file():
        raise HTTPException(status_code=404, detail="The stored Word template is unavailable.")
    return FileResponse(path, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", filename=document.original_filename)


@router.get("/templates/{template_id}/schema")
async def template_schema(template_id: str, user: User = Depends(require_permissions("ai.workspace.use", "knowledge.read")), session: AsyncSession = Depends(get_db_session)) -> dict:
    await _require_document_department(session, user)
    document, _ = await _template(session, user, template_id)
    document_type = _type_for(document.original_filename, document.source_key)
    row_fields = (["parameter", "specification", "result"] if document_type == "coa" else ["name", "cas_number", "ec_number", "concentration", "classification", "notes"])
    template_path = Path(get_settings().upload_storage_path) / document.stored_filename
    roles = await role_keys_for_user(session, user.id)
    default_rows = coa_parameter_rows(template_path) if document_type == "coa" else []
    return {"document_type": document_type, "fields": field_schema(document_type, template_path), "row_fields": row_fields, "default_rows": default_rows, "can_edit_filename": bool(roles.intersection({"owner", "super_admin", "department_admin"}))}


@router.get("/templates/{template_id}/excel-template")
async def excel_template(template_id: str, user: User = Depends(require_permissions("ai.workspace.use", "knowledge.read")), session: AsyncSession = Depends(get_db_session)) -> StreamingResponse:
    await _require_document_department(session, user)
    document, _ = await _template(session, user, template_id)
    document_type = _type_for(document.original_filename, document.source_key)
    workbook = Workbook()
    fields_sheet = workbook.active
    fields_sheet.title = "Fields"
    fields_sheet.append(["Field", "Value"])
    template_path = Path(get_settings().upload_storage_path) / document.stored_filename
    for item in field_schema(document_type, template_path):
        fields_sheet.append([item["label"], ""])
    rows_sheet = workbook.create_sheet("Parameters" if document_type == "coa" else "Composition")
    headers = ["Parameter", "Specification", "Result"] if document_type == "coa" else ["Name", "CAS Number", "EC Number", "Concentration", "Classification", "Notes"]
    rows_sheet.append(headers)
    if document_type == "coa":
        for row in coa_parameter_rows(template_path):
            rows_sheet.append([row["parameter"], "", ""])
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"{document_type.upper()}-input-template.xlsx"
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@router.post("/transcribe")
async def transcribe_draft_audio(
    audio_file: UploadFile = File(...),
    user: User = Depends(require_permissions("ai.workspace.use", "knowledge.read")),
    session: AsyncSession = Depends(get_db_session),
) -> dict:
    """Recheck the complete recording on Done; the audio is never stored."""
    await _require_document_department(session, user)
    settings = get_settings()
    if not settings.openai_api_key:
        raise HTTPException(status_code=503, detail="Professional voice transcription is not configured. The visible browser transcript can still be used.")
    content = await audio_file.read(15 * 1024 * 1024 + 1)
    if not content or len(content) > 15 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="The voice recording is empty or exceeds the 15 MB limit.")
    started = time.perf_counter()
    try:
        timeout = httpx.Timeout(settings.ai_request_timeout_seconds, connect=settings.ai_connect_timeout_seconds)
        async with httpx.AsyncClient(timeout=timeout) as client:
            response = await client.post(
                "https://api.openai.com/v1/audio/transcriptions",
                headers={"Authorization": f"Bearer {settings.openai_api_key}"},
                files={"file": (audio_file.filename or "rnd-draft.webm", content, audio_file.content_type or "audio/webm")},
                data={
                    "model": settings.openai_transcription_model,
                    "response_format": "json",
                    "prompt": "AROMAZEN professional COA and SDS dictation. Vocabulary: Name of Product, Product Code, Batch Number, Customer Name, Date of Manufacturing, Expiry Date, Quantity, Storage Condition, Appearance, pale yellowish liquid, Odour, Specific Gravity, Flash Point, Fire Point, Refractive Index, Tested By, Checked By. Preserve individually spoken letters exactly, especially F P versus S P. Preserve complete four-digit years such as 2028. Preserve verbal numeric ranges such as 0.85 to 1.15. Preserve corrections such as sorry, no, correct that to, and change that to.",
                },
            )
        if response.status_code >= 400:
            raise HTTPException(status_code=503, detail="The full voice recording could not be transcribed. Please review the visible transcript and try Done again.")
        payload = response.json()
        text = str(payload.get("text") or "").strip()
        if not text:
            raise HTTPException(status_code=422, detail="No clear speech was found in the recording.")
    except HTTPException:
        raise
    except (httpx.HTTPError, ValueError, json.JSONDecodeError) as exc:
        raise HTTPException(status_code=503, detail="The full voice recording could not be transcribed. Please review the visible transcript and try Done again.") from exc
    session.add(AIUsageEvent(
        organization_id=user.organization_id, user_id=user.id, department_id=user.department_id,
        operation="document_transcription", provider="openai", model=settings.openai_transcription_model,
        input_tokens=0, output_tokens=0, cost_usd=0,
        latency_ms=int((time.perf_counter() - started) * 1000), status="completed",
    ))
    await session.commit()
    return {"text": text}


@router.post("/draft-from-notes")
async def draft_from_notes(payload: DraftNotesRequest, user: User = Depends(require_permissions("ai.workspace.use", "knowledge.read")), session: AsyncSession = Depends(get_db_session)) -> dict:
    await _require_document_department(session, user)
    document, _ = await _template(session, user, payload.template_document_id)
    document_type = _type_for(document.original_filename, document.source_key)
    template_path = Path(get_settings().upload_storage_path) / document.stored_filename
    schema = field_schema(document_type, template_path)
    allowed_fields = {item["key"] for item in schema}
    fixed_parameters = coa_parameter_rows(template_path) if document_type == "coa" else []
    quick_updates = _quick_field_updates(payload.notes, document_type, allowed_fields)
    complex_markers = ("appearance", "odour", "odor", "specific gravity", "flash point", "fire point", "refractive index", "specification", "result", "composition", "ingredient", "cas", "concentration", "hazard", "precaution", "classification")
    needs_ai = not quick_updates or any(marker in payload.notes.lower() for marker in complex_markers)
    if not needs_ai:
        return {"field_updates": quick_updates, "row_updates": [], "unassigned_notes": "", "provider": "local", "model": "instant-field-mapper"}
    system = """You are a precise professional COA/SDS dictation editor. Return JSON only with keys field_updates, row_updates, and unassigned_notes. Input may contain a professional audio transcript and a browser transcript of the same speech; reconcile them as alternate evidence, do not treat their headings as values, and prefer the version that is complete and professionally plausible. Process notes in spoken order. When the speaker says 'sorry', 'no', 'I mean', 'correct that to', 'change that to', or otherwise revises a fact, replace the earlier value: the latest clear correction wins and old/new values must never be concatenated. Never invent. Correct only obvious speech-recognition artifacts using COA/SDS vocabulary, such as 'fail yellowish liquid' meaning 'pale yellowish liquid'. Omit genuinely unclear values. field_updates may use only supplied keys. For COA, 'name Rose' means product_name/Name of Product unless 'customer name' is explicitly spoken. Product codes and batch numbers must not absorb neighbouring fields. Preserve individually spoken code letters, full four-digit years, units, decimals, and numeric ranges. Expiry must not silently become the manufacturing year. Each value ends when another field or COA parameter is spoken. A COA parameter followed by one value normally updates result; if specification/result are named, map them exactly. Use only supplied fixed COA parameter names and never create or rename a parameter. For SDS, 'name Rose' means product_identifier. Treat deterministic updates as hints only: correct them when the complete sentence or a later correction provides a better value or boundary."""
    prompt = json.dumps({
        "document_type": document_type,
        "available_fields": [{"key": item["key"], "label": str(payload.field_labels.get(item["key"]) or item["label"])[:160]} for item in schema],
        "fixed_coa_parameters": [row["parameter"] for row in fixed_parameters],
        "current_fields": {key: str(value)[:4000] for key, value in payload.current_fields.items() if key in allowed_fields and str(value).strip()},
        "current_rows": payload.current_rows[:200],
        "deterministic_field_updates": quick_updates,
        "spoken_notes": payload.notes,
    }, ensure_ascii=False)
    started = time.perf_counter()
    provider = model = ""
    input_tokens = output_tokens = 0
    answer = ""
    try:
        runtime_settings = await provider_runtime_settings(session, user.organization_id)
        async for event in AIProviderRouter(runtime_settings).stream(system, prompt, payload.notes):
            if event.kind == "meta":
                provider, model = event.provider, event.model
            elif event.kind == "delta":
                answer += event.text
            elif event.kind == "usage":
                input_tokens, output_tokens = event.input_tokens, event.output_tokens
        parsed = _parse_ai_json(answer)
    except (ProviderError, ValueError, json.JSONDecodeError) as exc:
        raise HTTPException(status_code=503, detail="The AI Draft Assistant could not organize those notes. Please try again.") from exc
    updates = parsed.get("field_updates") if isinstance(parsed.get("field_updates"), dict) else {}
    field_updates = {str(key): _normalize_field_value(str(key), str(value)) for key, value in updates.items() if key in allowed_fields and str(value).strip()}
    # The complete utterance and its corrections take priority. The fast local
    # parser fills only values the AI did not confidently return.
    for key, value in quick_updates.items():
        field_updates.setdefault(key, value)
    raw_rows = parsed.get("row_updates") if isinstance(parsed.get("row_updates"), list) else []
    if document_type == "coa":
        allowed_parameters = {normalise(row["parameter"]): row["parameter"] for row in fixed_parameters}
        row_updates = []
        for row in raw_rows[:50]:
            if not isinstance(row, dict) or normalise(row.get("parameter")) not in allowed_parameters:
                continue
            parameter = allowed_parameters[normalise(row.get("parameter"))]
            row_updates.append({"parameter": parameter, "specification": _normalize_coa_value(parameter, "specification", row.get("specification") or ""), "result": _normalize_coa_value(parameter, "result", row.get("result") or "")})
    else:
        row_updates = [{str(key): str(value)[:2000] for key, value in row.items()} for row in raw_rows[:200] if isinstance(row, dict)]
    latency_ms = int((time.perf_counter() - started) * 1000)
    if provider and model:
        session.add(AIUsageEvent(organization_id=user.organization_id, user_id=user.id, department_id=user.department_id, operation="document_draft", provider=provider, model=model, input_tokens=input_tokens, output_tokens=output_tokens, cost_usd=estimate_cost(provider, model, input_tokens, output_tokens), latency_ms=latency_ms, status="completed"))
        await session.commit()
    return {"field_updates": field_updates, "row_updates": row_updates, "unassigned_notes": str(parsed.get("unassigned_notes") or "")[:4000], "provider": provider, "model": model}


@router.post("/generate")
async def generate(
    template_document_id: str = Form(...), document_type: str = Form(...), fields_json: str = Form("{}"), rows_json: str = Form("[]"),
    field_labels_json: str = Form("{}"), column_labels_json: str = Form("{}"), hidden_field_keys_json: str = Form("[]"), custom_fields_json: str = Form("[]"), output_filename: str | None = Form(None), excel_file: UploadFile | None = File(None), user: User = Depends(require_permissions("ai.workspace.use", "knowledge.read")), session: AsyncSession = Depends(get_db_session),
) -> dict:
    document_department_slug = await _require_document_department(session, user)
    document, _ = await _template(session, user, template_document_id)
    actual_type = _type_for(document.original_filename, document.source_key)
    if document_type not in {"coa", "sds"} or document_type != actual_type:
        raise HTTPException(status_code=422, detail="The selected document type does not match this template.")
    try:
        manual_fields = json.loads(fields_json)
        manual_rows = json.loads(rows_json)
        field_labels = json.loads(field_labels_json)
        column_labels = json.loads(column_labels_json)
        hidden_field_keys = json.loads(hidden_field_keys_json)
        custom_fields = json.loads(custom_fields_json)
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=422, detail="The form data is invalid.") from exc
    if not isinstance(manual_fields, dict) or not isinstance(manual_rows, list) or not isinstance(field_labels, dict) or not isinstance(column_labels, dict) or not isinstance(hidden_field_keys, list) or not isinstance(custom_fields, list):
        raise HTTPException(status_code=422, detail="The form data is invalid.")
    storage = Path(get_settings().upload_storage_path)
    excel_fields: dict[str, str] = {}
    excel_rows: list[dict[str, str]] = []
    temporary_excel: Path | None = None
    excel_content: bytes | None = None
    if excel_file:
        if Path(excel_file.filename or "").suffix.lower() != ".xlsx":
            raise HTTPException(status_code=422, detail="Please upload an XLSX Excel file.")
        max_excel_bytes = get_settings().max_excel_upload_size_mb * 1024 * 1024
        excel_content = await excel_file.read(max_excel_bytes + 1)
        if len(excel_content) > max_excel_bytes:
            raise HTTPException(status_code=413, detail=f"The Excel file exceeds the {get_settings().max_excel_upload_size_mb} MB limit.")
        temporary_excel = storage / f"{uuid.uuid4()}.xlsx"
        temporary_excel.write_bytes(excel_content)
        try:
            excel_fields, excel_rows = read_excel(temporary_excel, document_type)
        except Exception as exc:
            raise HTTPException(status_code=422, detail="The Excel file could not be read. Please use the downloadable input format.") from exc
        finally:
            temporary_excel.unlink(missing_ok=True)
    fields = {**excel_fields, **{str(k): str(v)[:4000] for k, v in manual_fields.items() if str(v).strip()}}
    if document_type == "coa":
        fields = {key: _normalize_field_value(key, value) for key, value in fields.items()}
    if document_type == "coa":
        combined: dict[str, dict[str, str]] = {normalise(row.get("parameter")): dict(row) for row in excel_rows if row.get("parameter")}
        for row in manual_rows:
            if not isinstance(row, dict) or not row.get("parameter"):
                continue
            key = normalise(row.get("parameter"))
            existing = combined.get(key, {"parameter": str(row.get("parameter")), "specification": "", "result": ""})
            for column in ("specification", "result"):
                if str(row.get(column) or "").strip():
                    existing[column] = str(row[column])
            combined[key] = existing
        rows = list(combined.values())
    else:
        rows = manual_rows if manual_rows else excel_rows
    rows = [{str(k): str(v)[:2000] for k, v in row.items()} for row in rows[:200] if isinstance(row, dict)]
    if document_type == "coa":
        rows = [{**row, "specification": _normalize_coa_value(row.get("parameter", ""), "specification", row.get("specification", "")), "result": _normalize_coa_value(row.get("parameter", ""), "result", row.get("result", ""))} for row in rows]
    template_path = storage / document.stored_filename
    if not template_path.is_file():
        raise HTTPException(status_code=404, detail="The stored Word template is unavailable.")
    generated_id = uuid.uuid4()
    product = fields.get("product_name") or fields.get("product_identifier") or document_type.upper()
    safe_product = re.sub(r"[^A-Za-z0-9_-]+", "-", product).strip("-")[:80] or document_type.upper()
    roles = await role_keys_for_user(session, user.id)
    can_edit_filename = bool(roles.intersection({"owner", "super_admin", "department_admin"}))
    if output_filename and not can_edit_filename:
        raise HTTPException(status_code=403, detail="Only administrators can edit the document filename.")
    requested_name = Path(output_filename or "").stem
    safe_requested_name = re.sub(r"[^A-Za-z0-9 _-]+", "", requested_name).strip()[:120]
    output_name = f"{safe_requested_name}.docx" if safe_requested_name else f"{safe_product}-{document_type.upper()}-DRAFT.docx"
    output_stored = organized_storage_name(
        "generated-documents",
        user.organization_id,
        output_name,
        category=document_type,
        identifier=generated_id,
    )
    (storage / output_stored).parent.mkdir(parents=True, exist_ok=True)
    safe_field_labels = {str(key): str(value)[:160] for key, value in field_labels.items() if str(value).strip()}
    safe_column_labels = {str(key): str(value)[:160] for key, value in column_labels.items() if str(value).strip()}
    allowed_field_keys = {item["key"] for item in field_schema(document_type, template_path)}
    safe_hidden_field_keys = {str(key) for key in hidden_field_keys if str(key) in allowed_field_keys}
    safe_custom_fields = [{"label": str(item.get("label", ""))[:160], "value": str(item.get("value", ""))[:4000]} for item in custom_fields[:50] if isinstance(item, dict) and str(item.get("label", "")).strip()]
    warnings = generate_docx(template_path, storage / output_stored, document_type, fields, rows, safe_field_labels, safe_column_labels, safe_hidden_field_keys, safe_custom_fields)
    if document_type == "coa" and fields.get("manufacturing_date") and fields.get("expiry_date"):
        try:
            manufacturing = datetime.strptime(fields["manufacturing_date"], "%d %B %Y")
            expiry = datetime.strptime(fields["expiry_date"], "%d %B %Y")
            if expiry <= manufacturing:
                warnings.append("Expiry Date is not later than Date of Manufacturing. Please correct it before issuing this draft.")
        except ValueError:
            warnings.append("One or more COA dates could not be validated. Please review them before issuing this draft.")
    generation = DocumentGeneration(id=generated_id, organization_id=user.organization_id, user_id=user.id, department_id=user.department_id, template_document_id=document.id, document_type=document_type, input_mode="mixed" if excel_file and manual_fields else "excel" if excel_file else "manual", output_stored_filename=output_stored, output_original_filename=output_name, warnings_json=warnings, status="draft")
    session.add(generation)
    session.add(AuditEvent(organization_id=user.organization_id, actor_user_id=user.id, action="document.generated", target_type="document_generation", target_id=str(generated_id), metadata_json={"document_type": document_type, "template_document_id": str(document.id), "input_mode": generation.input_mode, "warning_count": len(warnings)}))
    if excel_content is not None and excel_file is not None:
        await replace_department_uploads(session, user, document_department_slug, [DepartmentUpload(
            f"document-generator:{document_type}-data",
            excel_content,
            excel_file.filename or f"{document_type.upper()}_Input.xlsx",
            excel_file.content_type,
        )])
    else:
        await session.commit()
    return {"id": str(generated_id), "filename": output_name, "status": "draft", "warnings": warnings}


@router.get("/generations/{generation_id}/download")
async def download(generation_id: str, user: User = Depends(require_permissions("ai.workspace.use")), session: AsyncSession = Depends(get_db_session)) -> FileResponse:
    await _require_document_department(session, user)
    generation = await session.get(DocumentGeneration, generation_id)
    roles = await role_keys_for_user(session, user.id)
    if not generation or generation.organization_id != user.organization_id or (generation.user_id != user.id and not roles.intersection({"owner", "super_admin"})):
        raise HTTPException(status_code=404, detail="Generated document not found.")
    path = Path(get_settings().upload_storage_path) / generation.output_stored_filename
    if not path.is_file():
        raise HTTPException(status_code=404, detail="The generated file is unavailable.")
    return FileResponse(path, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", filename=generation.output_original_filename)


@router.get("/generations/{generation_id}/preview")
async def preview_generation(generation_id: str, user: User = Depends(require_permissions("ai.workspace.use")), session: AsyncSession = Depends(get_db_session)) -> StreamingResponse:
    await _require_document_department(session, user)
    generation = await session.get(DocumentGeneration, generation_id)
    roles = await role_keys_for_user(session, user.id)
    if not generation or generation.organization_id != user.organization_id or (generation.user_id != user.id and not roles.intersection({"owner", "super_admin"})):
        raise HTTPException(status_code=404, detail="Generated document not found.")
    path = Path(get_settings().upload_storage_path) / generation.output_stored_filename
    if not path.is_file():
        raise HTTPException(status_code=404, detail="The generated file is unavailable.")
    try:
        with TemporaryDirectory(prefix="coa-preview-") as directory:
            pdf_path = _convert_docx_to_pdf(path, Path(directory))
            pdf_bytes = pdf_path.read_bytes()
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail="PDF preview is unavailable because document conversion is not configured.") from exc
    pdf_name = f"{Path(generation.output_original_filename).stem}.pdf"
    return StreamingResponse(io.BytesIO(pdf_bytes), media_type="application/pdf", headers={"Content-Disposition": f'inline; filename="{pdf_name}"'})
