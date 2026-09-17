from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from app.modules.regulatory.engine import generate_regulatory_docx, _set_footer_metadata
from app.modules.regulatory.pictograms import pictogram_codes
from app.modules.regulatory.transport import resolve_transport
from app.modules.regulatory.reference_catalog import apply_raw_material_reference, un3082_technical_names

TEMPLATES = Path(__file__).resolve().parents[1] / 'app/templates/regulatory'


def test_single_and_multiple_pictograms_are_embedded_without_template_artwork(tmp_path):
    template = Document(TEMPLATES / 'sds.docx')
    for p in template.paragraphs:
        for image in p._p.xpath('.//w:drawing | .//w:pict'):
            image.getparent().remove(image)
    blank = tmp_path / 'no-artwork.docx'
    template.save(blank)
    for label, expected in [('Irritant', ['GHS07']), ('GHS07; GHS09', ['GHS07','GHS09']), ('Flame over circle', ['GHS03']), ('Health hazard, Corrosion', ['GHS05','GHS08'])]:
        output = tmp_path / 'sds.docx'
        generate_regulatory_docx(blank, output, 'sds', 'BLEND', 'X1', {'pictograms':label}, [{'name':'LINALOOL','concentration':'3'}])
        d = Document(output)
        assert [s._inline.docPr.get('descr').split()[0] for s in d.inline_shapes] == expected
        for shape in d.inline_shapes:
            rid = shape._inline.xpath('.//a:blip')[0].get(qn('r:embed'))
            assert d.part.related_parts[rid].blob.startswith(b'\x89PNG')
        # The reference shows only the bold "Pictograms:" label; the meaning is
        # carried by the artwork paragraph below, not by written symbol names.
        pict = next(p for p in d.paragraphs if p.text.startswith('Pictograms'))
        assert pict.text.strip() == 'Pictograms:'
        assert pict.runs[0].bold is True and pict.runs[0].text == 'Pictograms:'
        assert not any('Exclamation' in r.text or 'Environment' in r.text for r in pict.runs)
    assert pictogram_codes({'pictograms':'', 'hazard_statements':'H317, May cause an allergic skin reaction.'}) == ['GHS07']


def test_sds_typography_matches_reference_pdf(tmp_path):
    """Reference: Arial 10 body, bold labels + regular values, Calibri table data."""
    output = tmp_path / 'sds.docx'
    generate_regulatory_docx(
        TEMPLATES/'sds.docx', output, 'sds', 'Cedar and Sage', 'FS 12388',
        {'version': '2.1', 'revision_date': '17-09-2026'},
        [{'name':'HEDIONE','concentration':'20'}, {'name':'ISO E SUPER','concentration':'10'},
         {'name':'CARRYOPHELLENE  OXIDE SS','concentration':'10'}, {'name':'ISO BORNYL ACETATE','concentration':'10'},
         {'name':'BACDANOL','concentration':'6'}, {'name':'MUSK T','concentration':'5'},
         {'name':'VERDYL ACETATE','concentration':'5'}, {'name':'BENZYL  SALICYLATE','concentration':'5'},
         {'name':'FLORASOL','concentration':'5'}, {'name':'TERPINYL ACETATE','concentration':'5'},
         {'name':'BETA IONONE','concentration':'5'}, {'name':'DHM','concentration':'5'},
         {'name':'LINALOOL','concentration':'3'}, {'name':'VANILLIN','concentration':'3'},
         {'name':'LINALYL ACETATE','concentration':'3'}],
    )
    d = Document(output)
    labels = [p for p in d.paragraphs if p.text.strip().split('\t')[0].rstrip().endswith(':') and '\t' in p.text]
    checked = 0
    for p in labels:
        label, _, value = p.text.partition('\t')
        if not value.strip() or not p.runs:
            continue
        lead = next((r for r in p.runs if r.text.strip()), None)
        if lead is None or ':' not in lead.text:
            continue
        # bold=None renders bold under the master's Heading 2 style; only an
        # explicit False would break the bold-label look of the reference PDF.
        assert lead.bold is not False, f'label not bold: {p.text[:50]!r}'
        if len(p.runs) > 1:
            # bold=None also renders regular under the master's non-bold styles.
            assert any(r.bold is not True for r in p.runs[1:] if r.text.strip()), f'value not regular: {p.text[:50]!r}'
        checked += 1
    assert checked >= 6
    tables = d.tables
    data_runs = [r for row in tables[0].rows[1:].cells for r in row.paragraphs[0].runs if r.text.strip()] if False else [
        r for row in tables[0].rows[1:] for r in row.cells[0].paragraphs[0].runs if r.text.strip()
    ]
    assert data_runs and all(r.font.name == 'Calibri' for r in data_runs)
    header_runs = [r for row in tables[0].rows[:1] for r in row.cells[0].paragraphs[0].runs if r.text.strip()]
    assert header_runs and all(r.font.name == 'Arial' and r.bold for r in header_runs)


def test_sds_requested_tables_do_not_repeat_headers(tmp_path):
    output = tmp_path / 'sds.docx'
    generate_regulatory_docx(TEMPLATES/'sds.docx', output, 'sds', 'BLEND', 'X1', {}, [{'name':'MUSK T','concentration':'100'}])
    d = Document(output)
    for index in (0,3,5):
        assert not d.tables[index]._tbl.xpath('.//w:tblHeader')


def test_transport_source_membership_and_mixture_evidence_are_distinct():
    ingredients = [{'name':'OTBCHA','concentration':'5'}, {'name':'MUSK T','concentration':'5'}]
    before = deepcopy(ingredients)
    assert un3082_technical_names(ingredients) == ['MUSK T','OTBCHA']
    assert resolve_transport({}, ingredients) == 'Not determined'
    assert resolve_transport({'transport_un_number':'UN 3082'}, ingredients) == 'UN3082'
    assert resolve_transport({}, [{'name':'MUSK T','concentration':'100'}]) == 'UN3082'
    assert resolve_transport({'classification':'Aquatic Chronic 2: H411'}, ingredients) == 'UN3082'
    assert ingredients == before
    item = {'name':'ISO E SUPER','concentration':'3'}
    apply_raw_material_reference(item)
    assert resolve_transport({}, [item]) == 'UN3082'  # 10 x Chronic 1 reaches 25%.
    assert resolve_transport({'transport_un_number':'Not regulated'}, [item]) == 'Not regulated'


def test_ifra_amendment_fonts_footer_and_amount_column(tmp_path):
    output = tmp_path / 'amendment.docx'
    generate_regulatory_docx(TEMPLATES/'ifra-amendment.docx', output, 'ifra_amendment', 'BLEND', 'AB 123', {'version':'2.1','revision_date':'17-09-2026'}, [])
    d = Document(output)
    labels = [p for p in d.paragraphs if 'PRODUCT NAME' in p.text or 'PRODUCT CODE' in p.text]
    assert len(labels) == 2
    assert {(r.font.name,r.font.size.pt) for p in labels for r in p.runs if r.text.strip()} == {('Arial',9)}
    footer = ' '.join(d.sections[0].footer._element.xpath('.//w:t/text()'))
    assert '166A' in footer and 'B 105 106' not in footer
    assert '17-09-2026' in footer and '13-07-2026' not in footer
    assert 'Version: 2.1' in footer
    edges = set()
    for t in d.tables:
        for row in t.rows:
            assert len(row._tr.tc_lst) == 5
            cell = row.cells[-1]
            assert cell.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
            assert cell.paragraphs[0].paragraph_format.left_indent.pt == 0
            assert cell.text.strip() in {'NIL','% con'}
            edges.add(tuple(c.width.twips for c in row.cells))
    assert len(edges) == 1


def test_certificate_and_allergen_value_columns(tmp_path):
    for kind, filename in [('ifra_certificate','ifra-certificate.docx'),('allergen_report','allergen-report.docx')]:
        output = tmp_path / filename
        generate_regulatory_docx(TEMPLATES/filename, output, kind, 'BLEND', 'AB 123', {}, [])
        d = Document(output)
        tables = d.tables[:1] if kind == 'ifra_certificate' else d.tables
        for t in tables:
            for row in t.rows:
                p = row.cells[-1].paragraphs[0]
                assert p.alignment == WD_ALIGN_PARAGRAPH.CENTER
                assert p.paragraph_format.left_indent.pt == 0
                assert p.paragraph_format.first_line_indent.pt == 0
                assert not p.text.startswith((' ', '\t', '\n'))
        if kind == 'ifra_certificate':
            original = Document(TEMPLATES/filename)
            assert [r.cells[-1].text.strip() for r in original.tables[0].rows] == [r.cells[-1].text for r in d.tables[0].rows]


def test_footer_metadata_split_across_runs():
    d = Document()
    p = d.sections[0].footer.paragraphs[0]
    for value in ['Version: ', '0.0', '\tDate: ', '13-07-2026']:
        p.add_run(value)
    _set_footer_metadata(d, {'version':'3','revision_date':'17-09-2026'})
    assert p.text == 'Version: 3\tDate: 17-09-2026'
