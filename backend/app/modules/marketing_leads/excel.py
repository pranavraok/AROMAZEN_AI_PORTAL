from __future__ import annotations

import io
from datetime import datetime
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADERS = [
    "Sample reference",
    "Sample date",
    "Lead submitted",
    "Company / customer",
    "Contact person",
    "Phone number",
    "Email",
    "Country",
    "Region",
    "Product / category",
    "Expected quantity",
    "Lead source",
    "Priority",
    "Marketing employee",
    "Lead status",
    "First merchandising response",
    "Handled by",
    "Fragrance name",
    "Fragrance code",
    "Application",
    "Sample quantity",
    "Sample cost",
    "Sample outcome",
    "Sample remark",
    "Customer requirement",
    "Lead notes",
]


def _label(value: Any) -> str:
    return str(value or "").replace("_", " ").strip().title()


def _excel_value(value: Any) -> Any:
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.astimezone(ZoneInfo("Asia/Kolkata")).replace(tzinfo=None)
    return value


def build_merchandising_sample_requests_workbook(
    rows: list[dict[str, Any]],
    *,
    generated_at: datetime,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sample Requests"
    sheet.sheet_view.showGridLines = False
    last_column = get_column_letter(len(HEADERS))

    sheet.merge_cells(f"A1:{last_column}1")
    title = sheet["A1"]
    title.value = "Merchandising Sample Requests"
    title.font = Font(name="Arial", size=15, bold=True, color="173A33")
    title.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 27

    lead_count = len({row["lead_id"] for row in rows})
    sample_count = len({row["sample_id"] for row in rows})
    context = [
        ("A2", "Generated at", True),
        ("B2", _excel_value(generated_at), False),
        ("D2", "Lead handovers", True),
        ("E2", lead_count, False),
        ("G2", "Sample entries", True),
        ("H2", sample_count, False),
        ("J2", "Fragrance rows", True),
        ("K2", len(rows), False),
    ]
    for address, value, bold in context:
        cell = sheet[address]
        cell.value = value
        cell.font = Font(name="Arial", size=9, bold=bold, color="5F6F6B")
        cell.alignment = Alignment(vertical="center")
    sheet["B2"].number_format = "dd-mmm-yyyy hh:mm"
    sheet.row_dimensions[2].height = 20

    header_fill = PatternFill("solid", fgColor="2E5845")
    header_border = Border(
        right=Side(style="thin", color="FFFFFF"),
        bottom=Side(style="thin", color="21483A"),
    )
    for index, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(4, index, header)
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = header_border
    sheet.row_dimensions[4].height = 34

    keys = [
        "sample_reference",
        "sample_date",
        "lead_submitted_at",
        "company_name",
        "contact_person",
        "phone_number",
        "email",
        "country",
        "region",
        "product_interest",
        "expected_quantity",
        "lead_source",
        "priority",
        "marketing_employee",
        "lead_status",
        "merchandising_response_at",
        "handled_by",
        "fragrance_name",
        "fragrance_code",
        "application",
        "sample_quantity",
        "sample_cost",
        "sample_status",
        "sample_remark",
        "requirement",
        "lead_notes",
    ]
    status_keys = {"priority", "lead_status", "sample_status"}
    date_keys = {
        "sample_date",
        "lead_submitted_at",
        "merchandising_response_at",
    }
    pale_fill = PatternFill("solid", fgColor="F4F8F6")
    bottom_border = Border(bottom=Side(style="hair", color="DCE5E1"))
    for row_index, row in enumerate(rows, start=5):
        values = [
            _label(row.get(key)) if key in status_keys else _excel_value(row.get(key))
            for key in keys
        ]
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column_index, value)
            cell.font = Font(name="Arial", size=9, color="263A37")
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=column_index in {10, 18, 24, 25, 26},
            )
            cell.border = bottom_border
            if row_index % 2 == 0:
                cell.fill = pale_fill
        for key in date_keys:
            column_index = keys.index(key) + 1
            sheet.cell(row_index, column_index).number_format = (
                "dd-mmm-yyyy hh:mm" if key != "sample_date" else "dd-mmm-yyyy"
            )
        sheet.row_dimensions[row_index].height = 31

    widths = [
        17, 14, 19, 25, 20, 17, 27, 16, 17, 27, 18, 18, 12,
        22, 19, 21, 20, 27, 18, 22, 18, 15, 18, 30, 38, 32,
    ]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:{last_column}{max(4, len(rows) + 4)}"
    sheet.page_setup.orientation = "landscape"
    sheet.print_title_rows = "4:4"
    sheet.print_area = f"A1:{last_column}{max(4, len(rows) + 4)}"

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
