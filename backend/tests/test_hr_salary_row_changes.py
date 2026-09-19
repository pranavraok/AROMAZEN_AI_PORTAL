import sys
import tempfile
import types
import unittest
from pathlib import Path

# routes.py imports pdfplumber (used only for PDF text checks); the local test
# environment may not have it, and these tests never touch PDF extraction.
if "pdfplumber" not in sys.modules:
    sys.modules["pdfplumber"] = types.ModuleType("pdfplumber")

from docx import Document

from app.modules.hr_letters import routes as R
from app.modules.hr_letters.routes import AddedSalaryRow, LetterRequest


def _special_increment_document():
    return Document(R.ASSET_ROOT / R.TEMPLATE_FILES["special_increment"])


def _salary_rows(document):
    rows = []
    for table in R._iter_salary_tables(document):
        for row in table.rows:
            keys = R._row_salary_keys(row)
            if keys:
                rows.append((row, keys))
    return rows


class SalaryRowChangesTests(unittest.TestCase):
    def test_request_models_accept_row_changes(self) -> None:
        payload = LetterRequest(
            template_key="special_increment",
            fields={"employee_name": "Test"},
            deleted_salary_rows=["pf"],
            added_salary_rows=[AddedSalaryRow(key="food_allowance", label="Food Allowance")],
        )
        self.assertEqual(payload.deleted_salary_rows, ["pf"])
        self.assertEqual(payload.added_salary_rows[0].label, "Food Allowance")

    def test_added_row_clones_formatting_and_plants_tokens(self) -> None:
        document = _special_increment_document()
        before = _salary_rows(document)
        R._apply_salary_row_changes(
            document,
            [],
            [{"key": "food_allowance", "label": "Food Allowance"}],
        )
        after = _salary_rows(document)
        self.assertEqual(len(after), len(before) + 1)
        new_row, keys = after[-1]
        self.assertIn("food_allowance", keys)
        self.assertEqual(new_row.cells[0].text.strip(), "Food Allowance")
        # Cloned row keeps the source row's columns, re-keyed to the new row.
        source_columns = R._cell_salary_columns(before[0][0].cells[1])
        self.assertEqual(R._cell_salary_columns(new_row.cells[1]), source_columns)
        self.assertIn("{{salary_food_allowance_existing}}", new_row.cells[1].text)
        self.assertIn("{{salary_food_allowance_revised}}", new_row.cells[2].text)

    def test_nested_annexure_table_is_reached(self) -> None:
        document = _special_increment_document()
        top_level_keys = set()
        for table in document.tables:
            for row in table.rows:
                top_level_keys |= R._row_salary_keys(row)
        self.assertEqual(top_level_keys, set(), "annexure must be nested for this regression to matter")
        self.assertGreater(len(_salary_rows(document)), 0)

    def test_deleted_row_is_removed(self) -> None:
        document = _special_increment_document()
        before = _salary_rows(document)
        target = "pf"
        self.assertTrue(any(target in keys for _, keys in before), "template must contain the row to delete")
        R._apply_salary_row_changes(document, [target], [])
        after = _salary_rows(document)
        self.assertEqual(len(after), len(before) - 1)
        self.assertFalse(any(target in keys for _, keys in after))

    def test_invalid_keys_and_labels_are_ignored(self) -> None:
        document = _special_increment_document()
        before = _salary_rows(document)
        R._apply_salary_row_changes(
            document,
            ["NOT A KEY!!"],
            [{"key": "Bad Key!", "label": ""}, {"key": "ok_key", "label": "  "}],
        )
        self.assertEqual(len(_salary_rows(document)), len(before))

    def test_no_changes_is_a_noop(self) -> None:
        document = _special_increment_document()
        before = _salary_rows(document)
        R._apply_salary_row_changes(document, [], [])
        self.assertEqual(len(_salary_rows(document)), len(before))

    def test_fill_docx_applies_row_changes_end_to_end(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            docx_path = R._fill_docx(
                "special_increment",
                {"employee_name": "Test Employee"},
                Path(temporary),
                deleted_salary_rows=["pf"],
                added_salary_rows=[{"key": "food_allowance", "label": "Food Allowance"}],
            )
            document = Document(docx_path)
            texts = [cell.text for table in R._iter_salary_tables(document) for row in table.rows for cell in row.cells]
            self.assertFalse(any(R._row_salary_keys(row) & {"pf"} for table in R._iter_salary_tables(document) for row in table.rows))
            self.assertTrue(any(text.strip() == "Food Allowance" for text in texts))


if __name__ == "__main__":
    unittest.main()
