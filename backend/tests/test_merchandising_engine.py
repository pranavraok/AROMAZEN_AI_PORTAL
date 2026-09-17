from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from app.modules.merchandising.engine import DOCUMENT_TYPES, extract_fields, generate_document, missing_required


ASSETS = Path(__file__).resolve().parents[1] / "app" / "assets" / "merchandising"
FILES = {
    "packing_list": "packing-list.xlsx",
    "hazardous_request": "hazardous-cargo-request.xlsx",
    "imo_declaration": "imo-dangerous-goods-declaration.xlsx",
    "multimodal_form": "multimodal-dangerous-goods-form.docx",
}


def _merged_fields() -> dict[str, str]:
    values: dict[str, str] = {}
    for document_type in DOCUMENT_TYPES:
        filename = FILES[document_type]
        detected, extracted = extract_fields((ASSETS / filename).read_bytes(), filename)
        assert detected == document_type
        for key, value in extracted.items():
            values.setdefault(key, value)
    return values


def test_supplied_merchandising_masters_are_detected_and_cover_required_fields() -> None:
    fields = _merged_fields()
    assert fields["product_name"] == "CHERRY BOOSTER FP 12934"
    assert fields["un_number"].replace(" ", "") == "3082"
    assert fields["imo_class"] == "9"
    assert fields["gross_mass"]
    assert missing_required(fields) == []


def test_generates_all_four_documents_from_one_record(tmp_path: Path) -> None:
    fields = _merged_fields()
    outputs: dict[str, Path] = {}
    for document_type, filename in FILES.items():
        extension = ".docx" if document_type == "multimodal_form" else ".xlsx"
        output = tmp_path / f"{document_type}{extension}"
        generate_document(ASSETS / filename, output, document_type, fields)
        assert output.is_file() and output.stat().st_size > 5_000
        outputs[document_type] = output

    packing = load_workbook(outputs["packing_list"]).active
    assert packing["E28"].value == fields["product_name"]
    assert str(packing["J23"].value) == fields["gross_mass"]

    hazardous = load_workbook(outputs["hazardous_request"]).active
    assert str(hazardous["C13"].value).replace(" ", "") == "3082"
    assert hazardous["C14"].value == "9"

    imo = load_workbook(outputs["imo_declaration"]).active
    assert "3082" in str(imo["B33"].value)
    assert imo.print_area == "'Sheet1'!$A$1:$H$64"

    multimodal = Document(outputs["multimodal_form"])
    text = "\n".join(cell.text for row in multimodal.tables[0].rows for cell in row.cells)
    assert "3082" in text
    assert fields["product_name"] not in text or fields["proper_shipping_name"] in text

