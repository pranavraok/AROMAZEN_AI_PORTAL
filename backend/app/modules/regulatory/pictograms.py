"""Embed official GHS artwork independently of uploaded template drawings."""
from pathlib import Path
import re

from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

ASSETS = Path(__file__).resolve().parents[2] / "templates/regulatory/pictograms"
SYMBOLS = {
    "GHS01": ("Exploding bomb", ("explosive", "exploding bomb")),
    "GHS02": ("Flame", ("flammable", "flame")),
    "GHS03": ("Flame over circle", ("oxidizer", "oxidising", "oxidizing", "flame over circle")),
    "GHS04": ("Gas cylinder", ("gas cylinder", "gas under pressure")),
    "GHS05": ("Corrosion", ("corrosion", "corrosive")),
    "GHS06": ("Skull and crossbones", ("skull", "acute toxicity")),
    "GHS07": ("Exclamation mark", ("irritant", "exclamation", "harmful")),
    "GHS08": ("Health hazard", ("health hazard", "carcinogen", "respiratory sensitizer")),
    "GHS09": ("Environment", ("environment", "aquatic")),
}
HAZARDS = {
    "GHS02": {"H220", "H221", "H222", "H223", "H224", "H225", "H226", "H228", "H250", "H251", "H252", "H260", "H261"},
    "GHS03": {"H270", "H271", "H272"},
    "GHS04": {"H280", "H281"},
    "GHS05": {"H290", "H314", "H318"},
    "GHS06": {"H300", "H301", "H310", "H311", "H330", "H331"},
    "GHS07": {"H302", "H312", "H315", "H317", "H319", "H332", "H335", "H336"},
    "GHS08": {"H304", "H334", "H340", "H341", "H350", "H351", "H360", "H361", "H370", "H371", "H372", "H373"},
    "GHS09": {"H400", "H410", "H411"},
}


def pictogram_codes(label: dict[str, str]) -> list[str]:
    stated = str(label.get("pictograms") or "").strip()
    codes = set(re.findall(r"GHS\s*0?([1-9])\b", stated, re.I))
    selected = {f"GHS0{code}" for code in codes}
    for token in re.split(r"[,;\n/]+", stated.lower()):
        matches = [(len(alias), code) for code, (_, aliases) in SYMBOLS.items() for alias in aliases if alias in token]
        if matches:
            # "Flame over circle" must not also select "Flame".
            selected.add(max(matches)[1])
    if not selected and stated.lower() not in {"none", "not applicable", "no pictograms"}:
        hazards = set(re.findall(r"\bH\d{3}", f"{label.get('classification', '')} {label.get('hazard_statements', '')}"))
        selected = {code for code, values in HAZARDS.items() if hazards & values}
        if "GHS06" in selected:
            selected.discard("GHS07")
        if "GHS05" in selected and not hazards & {"H302", "H312", "H317", "H332", "H335", "H336"}:
            selected.discard("GHS07")
    return sorted(selected)


def fill_pictograms(document, label: dict[str, str]) -> None:
    paragraphs = list(document.paragraphs)
    start = next((i for i, p in enumerate(paragraphs) if re.match(r"^\s*Pictograms", p.text, re.I)), None)
    if start is None:
        return
    end = next((i for i in range(start + 1, len(paragraphs)) if re.match(r"^\s*(?:2\.3\s+)?Other hazards", paragraphs[i].text, re.I)), start + 1)
    for paragraph in paragraphs[start + 1:end]:
        if not paragraph._p.xpath(".//w:sectPr"):
            paragraph._p.getparent().remove(paragraph._p)
    codes = pictogram_codes(label)
    paragraph = paragraphs[start]
    # The approved reference shows only the bold label; the artwork below
    # carries the meaning, so no symbol names are written out.
    paragraph.text = "Pictograms:"
    paragraph.paragraph_format.keep_with_next = bool(codes)
    if paragraph.runs:
        paragraph.runs[0].text = "Pictograms:"
        paragraph.runs[0].bold = True
        for extra in paragraph.runs[1:]:
            extra.text = ""
    else:
        label_run = paragraph.add_run("Pictograms:")
        label_run.bold = True
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(10)
    if not codes:
        return
    node = OxmlElement("w:p")
    paragraph._p.addnext(node)
    pictures = Paragraph(node, paragraph._parent)
    pictures.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pictures.paragraph_format.left_indent = Inches(2.15)
    pictures.paragraph_format.space_after = Pt(6)
    for code in codes:
        shape = pictures.add_run().add_picture(str(ASSETS / f"{code}.png"), width=Inches(0.62))
        shape._inline.docPr.set("descr", f"{code} {SYMBOLS[code][0]}")
        pictures.add_run(" ")
